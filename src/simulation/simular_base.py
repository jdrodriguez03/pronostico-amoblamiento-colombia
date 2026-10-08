"""
Construccion de la base de datos del Proyecto Integrador ML I.

Variable objetivo: indice de ventas reales (base 2019 = 100) por
departamento (7 dominios EMC) x linea del hogar (8 lineas) x mes.

    Indice(d, l, t) = L(l, t) * F(d, t) * eta(d, l, t)

    L   : indice nacional real de la linea l           (EMC, anexo nacional, hoja 2.2)
    F   : factor regional del grupo CIIU 4741-4759       (EMC, anexo departamental, hoja 3.2)
    eta : componente especifico linea x departamento, AR(1), unica parte simulada

Uso (desde la raiz del repositorio):
    python src/simulation/simular_base.py
        lee data/raw/, escribe las bases en data/interim/ y la verificacion en results/
    python src/simulation/simular_base.py <anexo_nacional.xlsx> <anexo_departamental.xlsx> <carpeta_salida>
"""
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

# Rutas por defecto, relativas a la raiz del repositorio
RAIZ = Path(__file__).resolve().parents[2]
ANEXO_NACIONAL = RAIZ / "data" / "raw" / "anex-EMC-ComercioMinorista-jul2026.xlsx"
ANEXO_DEPARTAMENTAL = RAIZ / "data" / "raw" / "anex-EMC-ComercioalpormenorDep-jul2026.xlsx"
SALIDA = RAIZ / "data" / "interim"
VERIFICACION = RAIZ / "results" / "verificacion.md"

# ---------------------------------------------------------------------------
# Parametros (todos fijados antes de modelar; ver informe, seccion 6.3)
# ---------------------------------------------------------------------------
SEMILLA = 42
RHO = 0.6                                   # persistencia de eta (autocorrelacion observada 0,5-0,66)
ESCENARIOS = {"bajo": 0.02, "central": 0.07, "alto": 0.13}   # desviacion de eta
LINEAS = {                                  # codigo EMC -> nombre corto
    4: "Prendas de vestir y textiles",
    8: "Electrodomesticos y muebles",
    9: "Articulos y utensilios domesticos",
    10: "Aseo del hogar",
    11: "Informatica y telecomunicaciones",
    12: "Sonido y video",
    14: "Ferreteria, vidrios y pinturas",
    15: "Otras mercancias de uso domestico",
}
DEPTOS = ["Antioquia", "Atlantico", "Bogota", "Cundinamarca",
          "Santander", "Valle del Cauca", "Otros departamentos"]
GRUPO_HOGAR = "4741"                        # grupo CIIU 4741-4759 (incluye 4754)
FILA_BASE_PESOS = 78                        # julio 2025 (mes base de las contribuciones publicadas)


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return np.nan


def hoja(wb, nombre):
    for s in wb.sheetnames:
        if s.strip() == nombre:
            return list(wb[s].iter_rows(values_only=True))
    raise KeyError(nombre)


def filas_datos(rows, desde):
    return [r for r in rows[desde:] if np.isfinite(num(r[3]))]


# ---------------------------------------------------------------------------
# 1. Lineas nacionales (L)
# ---------------------------------------------------------------------------
def leer_lineas(ruta):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    rows = hoja(wb, "2.2")
    hdr = rows[7]
    col = {}
    for j, h in enumerate(hdr):
        if h is None:
            continue
        txt = str(h).strip()
        cod = txt.split(".")[0].strip()
        if cod.isdigit() and int(cod) in LINEAS:
            col[int(cod)] = j
    datos = filas_datos(rows, 8)
    fechas = pd.date_range("2019-01-01", periods=len(datos), freq="MS")
    L = np.array([[num(r[col[c]]) for c in LINEAS] for r in datos])
    return fechas, L, wb


# ---------------------------------------------------------------------------
# 2. Factor regional (F) y pesos departamentales
# ---------------------------------------------------------------------------
def leer_departamentos(ruta, n_meses):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)

    # Peso de cada departamento en el comercio total = contribucion / variacion (hoja 1.1)
    peso_total = []
    for r in hoja(wb, "1.1")[10:17]:
        v = [x for x in r if x not in (None, "")]
        peso_total.append(num(v[4]) / num(v[3]))
    peso_total = np.array(peso_total)

    # Participacion del grupo 4741-4759 dentro de cada departamento (hoja 2.1)
    rows = hoja(wb, "2.1")
    hdr, sub = rows[7], rows[8]
    cols, cur = [], None
    for j, h in enumerate(hdr):
        if h:
            cur = str(h).strip()
        if cur and sub[j] and str(sub[j]).strip() == "Real" and cur not in [c[0] for c in cols]:
            cols.append((cur, j))
    cols = cols[:7]
    fila_g = next(r for r in rows if r and any(GRUPO_HOGAR in str(x) for x in r if x))
    part_grupo = np.array([num(fila_g[j + 1]) / num(fila_g[j]) for _, j in cols])

    # Indices reales del grupo por departamento (hoja 3.2)
    rows = hoja(wb, "3.2")
    gcols = [j for j, h in enumerate(rows[8]) if h and GRUPO_HOGAR in str(h)]
    datos = filas_datos(rows, 9)
    G = np.array([[num(r[j]) for j in gcols] for r in datos])
    assert G.shape == (n_meses, 7), "Las series departamentales no coinciden con las nacionales"

    # Pesos del grupo a julio 2025, llevados a base 2019
    s = peso_total * part_grupo
    s = s / s.sum()
    w = s / G[FILA_BASE_PESOS]
    w = w / w.sum()

    R = G @ w                       # agregado nacional del grupo
    F = G / R[:, None]              # factor regional
    return F, G, w, R


# ---------------------------------------------------------------------------
# 3. Componente simulado (eta) y construccion
# ---------------------------------------------------------------------------
def simular(L, F, w, sigma, rng):
    T, nl = L.shape
    nd = F.shape[1]
    e = np.zeros((T, nd, nl))
    innov = sigma * np.sqrt(1 - RHO ** 2)           # mantiene la desviacion estacionaria = sigma
    e[0] = rng.normal(0, sigma, (nd, nl))
    for t in range(1, T):
        e[t] = RHO * e[t - 1] + rng.normal(0, innov, (nd, nl))
    eta = np.exp(e - sigma ** 2 / 2)
    Y = L[:, None, :] * F[:, :, None] * eta
    agregado = np.einsum("d,tdl->tl", w, Y)
    Y = Y * (L / agregado)[:, None, :]              # conservacion exacta de la serie nacional
    return Y


def a_tabla(Y, fechas, escenario):
    T, nd, nl = Y.shape
    cods = list(LINEAS)
    reg = []
    for t in range(T):
        for d in range(nd):
            for l in range(nl):
                reg.append((fechas[t], fechas[t].year, fechas[t].month, DEPTOS[d],
                            cods[l], LINEAS[cods[l]], int(cods[l] == 8), escenario, Y[t, d, l]))
    return pd.DataFrame(reg, columns=["fecha", "anio", "mes", "departamento", "linea_codigo",
                                      "linea", "es_muebles", "escenario", "indice_ventas_real"])


# ---------------------------------------------------------------------------
# 4. Verificaciones (rubrica 7.1)
# ---------------------------------------------------------------------------
def verificar(Y, L, F, G, w, R, C474_475, sigma, nombre):
    out = [f"## Escenario {nombre} (desviacion de eta = {sigma:.0%})"]
    n = Y.size
    out.append(f"1. Integridad: {n} registros | vacios {np.isnan(Y).sum()} | no positivos {(Y <= 0).sum()}")
    agg = np.einsum("d,tdl->tl", w, Y)
    out.append(f"2. Conservacion nacional: error maximo {np.max(np.abs(agg / L - 1)) * 100:.1e} %")
    r = np.log(Y / (L[:, None, :] * F[:, :, None]))
    ac = np.mean([np.corrcoef(r[1:, d, l], r[:-1, d, l])[0, 1]
                  for d in range(r.shape[1]) for l in range(r.shape[2])])
    out.append(f"3. Ruido recuperado: desviacion {r.std() * 100:.1f} % (objetivo {sigma * 100:.0f} %) | "
               f"autocorrelacion {ac:.2f} (objetivo {RHO})")
    i8 = list(LINEAS).index(8)
    cr = [np.corrcoef(Y[:, d, i8], G[:, d])[0, 1] for d in range(7)]
    out.append(f"4. Fidelidad regional (muebles vs grupo real del depto): correlacion {min(cr):.2f} a {max(cr):.2f}")
    mu = pd.Series(np.einsum("d,td->t", w, Y[:, :, i8]))
    out.append(f"5. Patrones: muebles nacional abr-2020 = {mu[15]:.1f} (real {L[15, i8]:.1f})")
    return "\n".join(out)


def main(nac, dep, salida):
    salida = Path(salida)
    salida.mkdir(parents=True, exist_ok=True)
    fechas, L, wb_nac = leer_lineas(nac)
    F, G, w, R = leer_departamentos(dep, len(fechas))

    # Verificacion contable de los pesos contra las actividades 474 y 475 (anexo nacional, hoja 2.4)
    rows = hoja(wb_nac, "2.4")
    j474 = next(j for j, h in enumerate(rows[7]) if h and str(h).strip().startswith("474"))
    j475 = next(j for j, h in enumerate(rows[7]) if h and str(h).strip().startswith("475"))
    C = np.array([[num(r[j474]), num(r[j475])] for r in filas_datos(rows, 8)])
    b = np.linalg.lstsq(C, R, rcond=None)[0]
    err_pesos = np.max(np.abs(C @ b - R) / R) * 100

    informe = ["# Verificacion de la base construida", "",
               f"Semilla: {SEMILLA} | Meses: {len(fechas)} ({fechas[0]:%Y-%m} a {fechas[-1]:%Y-%m}) | "
               f"Departamentos: 7 | Lineas: {len(LINEAS)}", "",
               "Pesos departamentales (base 2019): " +
               ", ".join(f"{d} {x:.1%}" for d, x in zip(DEPTOS, w)),
               f"Verificacion contable de pesos vs actividades 474/475: error maximo {err_pesos:.1e} %", ""]

    rng = np.random.default_rng(SEMILLA)
    tablas = []
    for nombre, sigma in ESCENARIOS.items():
        Y = simular(L, F, w, sigma, rng)
        tablas.append(a_tabla(Y, fechas, nombre))
        informe += [verificar(Y, L, F, G, w, R, C, sigma, nombre), ""]

    base = pd.concat(tablas, ignore_index=True)
    for nombre in ESCENARIOS:
        base[base.escenario == nombre].drop(columns="escenario").to_csv(
            salida / f"base_{nombre}.csv", index=False, float_format="%.4f")

    # Componentes reales (para auditoria; NO se usan como predictores)
    comp = pd.DataFrame([(fechas[t], DEPTOS[d], w[d], F[t, d], G[t, d], R[t])
                         for t in range(len(fechas)) for d in range(7)],
                        columns=["fecha", "departamento", "peso_base2019", "factor_regional",
                                 "indice_grupo_depto", "indice_grupo_nacional"])
    comp.to_csv(salida / "componentes_regionales.csv", index=False, float_format="%.6f")
    lin = pd.DataFrame(L, columns=[f"L_{c}" for c in LINEAS])
    lin.insert(0, "fecha", fechas)
    lin.to_csv(salida / "lineas_nacionales.csv", index=False, float_format="%.6f")

    VERIFICACION.parent.mkdir(parents=True, exist_ok=True)
    VERIFICACION.write_text("\n".join(informe), encoding="utf-8")
    print("\n".join(informe))


if __name__ == "__main__":
    main(*(sys.argv[1:4] or [ANEXO_NACIONAL, ANEXO_DEPARTAMENTAL, SALIDA]))
