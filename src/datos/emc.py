"""Lectura de las series de ventas reales de la Encuesta Mensual de Comercio."""

from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from src.config import (
    DEPARTAMENTOS,
    FECHA_INICIO,
    LINEAS,
    N_MESES,
    PESOS_DEPTO,
    RUTA_RAW,
)

ANEXO_NACIONAL = RUTA_RAW / "anex-EMC-ComercioMinorista-jul2026.xlsx"
ANEXO_DEPARTAMENTAL = RUTA_RAW / "anex-EMC-ComercioalpormenorDep-jul2026.xlsx"
GRUPO_CIIU_HOGAR = "4741"


def _leer_hoja(ruta: Path, nombre: str) -> list[tuple]:
    """Carga una hoja del libro como filas, respetando los valores guardados en Excel."""
    libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        hoja = next((hoja for hoja in libro.sheetnames if hoja.strip() == nombre), None)
        if hoja is None:
            raise KeyError(f"No existe la hoja {nombre!r} en {ruta.name}")
        return list(libro[hoja].iter_rows(values_only=True))
    finally:
        libro.close()


def _numero(valor: object) -> float:
    """Convierte una celda numérica de Excel a float; valores vacíos quedan como NaN."""
    try:
        return float(valor)
    except (TypeError, ValueError):
        return np.nan


def _fechas_mensuales(cantidad: int) -> pd.DatetimeIndex:
    fechas = pd.date_range(FECHA_INICIO, periods=cantidad, freq="MS")
    if cantidad != N_MESES:
        raise ValueError(f"Se esperaban {N_MESES} meses, se encontraron {cantidad}")
    return fechas


def leer_lineas_nacional(ruta: Path = ANEXO_NACIONAL) -> pd.DataFrame:
    """Lee las ocho líneas nacionales de ventas reales de la hoja 2.2.

    Devuelve una fila por mes y línea, con `fecha` (datetime), `linea` (str) y
    `L` (float).
    """
    filas = _leer_hoja(ruta, "2.2")
    encabezado = filas[7]
    columnas = {}
    for posicion, celda in enumerate(encabezado):
        if celda is None:
            continue
        codigo = str(celda).strip().split(".", maxsplit=1)[0].strip()
        if codigo.isdigit() and int(codigo) in LINEAS:
            columnas[int(codigo)] = posicion

    if set(columnas) != set(LINEAS):
        faltantes = sorted(set(LINEAS) - set(columnas))
        raise ValueError(f"No se encontraron todas las líneas EMC en la hoja 2.2: {faltantes}")

    registros = [fila for fila in filas[8:] if len(fila) > 3 and np.isfinite(_numero(fila[3]))]
    fechas = _fechas_mensuales(len(registros))
    largos = []
    for codigo, nombre in LINEAS.items():
        largos.append(
            pd.DataFrame(
                {
                    "fecha": fechas,
                    "linea": nombre,
                    "L": [_numero(fila[columnas[codigo]]) for fila in registros],
                }
            )
        )
    return pd.concat(largos, ignore_index=True)[["fecha", "linea", "L"]]


def leer_departamental(ruta: Path = ANEXO_DEPARTAMENTAL) -> pd.DataFrame:
    """Lee en la hoja 3.2 el índice real del grupo CIIU 4741–4759.

    Devuelve una fila por mes y dominio, con `fecha` (datetime),
    `departamento` (str) e `indice_depto` (float).
    """
    filas = _leer_hoja(ruta, "3.2")
    encabezados = filas[8]
    columnas_grupo = [
        posicion
        for posicion, celda in enumerate(encabezados)
        if celda is not None and GRUPO_CIIU_HOGAR in str(celda)
    ]
    if len(columnas_grupo) != len(DEPARTAMENTOS):
        raise ValueError(
            "La hoja 3.2 debe contener una columna del grupo CIIU "
            f"{GRUPO_CIIU_HOGAR} por cada dominio; se encontraron {len(columnas_grupo)}"
        )

    registros = [
        fila
        for fila in filas[9:]
        if all(
            posicion < len(fila) and np.isfinite(_numero(fila[posicion]))
            for posicion in columnas_grupo
        )
    ]
    fechas = _fechas_mensuales(len(registros))
    return pd.DataFrame(
        [
            {
                "fecha": fecha,
                "departamento": departamento,
                "indice_depto": _numero(fila[posicion]),
            }
            for posicion, departamento in zip(columnas_grupo, DEPARTAMENTOS)
            for fecha, fila in zip(fechas, registros)
        ],
        columns=["fecha", "departamento", "indice_depto"],
    )


def factor_regional(indices: pd.DataFrame | None = None) -> pd.DataFrame:
    """Normaliza cada índice departamental por el agregado ponderado del mes.

    Usa directamente los pesos base 2019 de `config.PESOS_DEPTO`. Devuelve las
    columnas `fecha`, `departamento` y `factor_regional`.
    """
    if indices is None:
        indices = leer_departamental()
    requeridas = {"fecha", "departamento", "indice_depto"}
    if not requeridas.issubset(indices.columns):
        raise ValueError(f"Faltan columnas requeridas: {sorted(requeridas - set(indices.columns))}")

    datos = indices[["fecha", "departamento", "indice_depto"]].copy()
    datos["peso"] = datos["departamento"].map(PESOS_DEPTO)
    if datos["peso"].isna().any():
        dominios = sorted(datos.loc[datos["peso"].isna(), "departamento"].unique())
        raise ValueError(f"Departamentos sin peso configurado: {dominios}")
    agregado = (datos["indice_depto"] * datos["peso"]).groupby(datos["fecha"]).transform("sum")
    if agregado.isna().any() or (agregado == 0).any():
        raise ValueError("El agregado ponderado debe ser finito y distinto de cero en cada fecha")
    datos["factor_regional"] = datos["indice_depto"] / agregado
    return datos[["fecha", "departamento", "factor_regional"]]
